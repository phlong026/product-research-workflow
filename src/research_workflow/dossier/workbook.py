from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from . import sections
from .models import ResearchDossierInput

NOT_PROVIDED = "本次运行未提供"
WORKBOOK_FILENAME = "research_workbook.xlsx"
WORKBOOK_SKIPPED_NOTE = (
    "综合 Excel 工作簿未生成:缺少 openpyxl(安装 optional extra: xlsx)"
)
DECISION_NOTE = "最终商业决定由人工填写,系统不代填"
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_FIXED_DATETIME = datetime(2000, 1, 1)


class WorkbookDependencyError(RuntimeError):
    """openpyxl 未安装(optional extra `xlsx`)。"""


def _import_openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - 由测试通过 monkeypatch 覆盖
        raise WorkbookDependencyError(
            "缺少 openpyxl,无法生成综合 Excel 工作簿;"
            "请安装 optional extra: pip install 'product-research-workflow[xlsx]'"
        ) from exc
    return openpyxl


def openpyxl_available() -> bool:
    try:
        _import_openpyxl()
    except WorkbookDependencyError:
        return False
    return True


@dataclass(frozen=True, slots=True)
class _Col:
    header: str
    get: Callable[[Any], object]
    kind: str = "text"  # text | int | num | pct(源数据为百分点,写成分数 + 0.00% 格式)


def _attr(name: str) -> Callable[[Any], object]:
    return lambda row: getattr(row, name)


def _c(header: str, name: str, kind: str = "text") -> _Col:
    return _Col(header, _attr(name), kind)


_MARKET = [
    _c("月份", "month"),
    _c("总销量", "total_sales", "int"),
    _c("ASIN数", "asin_count", "int"),
    _c("环比", "mom_pct", "pct"),
    _c("同比", "yoy_pct", "pct"),
]
_REGION = [
    _c("月份", "month"),
    _c("卖家地区", "seller_region"),
    _c("销量", "sales", "int"),
    _c("ASIN数", "asin_count", "int"),
    _c("销量份额", "sales_share_pct", "pct"),
    _c("ASIN份额", "asin_share_pct", "pct"),
    _c("销量环比", "sales_mom_pct", "pct"),
    _c("销量同比", "sales_yoy_pct", "pct"),
    _c("份额环比(pp)", "share_mom_pp", "num"),
    _c("份额同比(pp)", "share_yoy_pp", "num"),
    _c("覆盖提示", "coverage_warning"),
    _c("待验证假设", "cause_hypothesis"),
]
_REGION_SEGMENT = [
    _c("月份", "month"),
    _c("卖家地区", "seller_region"),
    _c("分类", "segment"),
    _c("销量", "sales", "int"),
    _c("ASIN数", "asin_count", "int"),
    _c("占当月销量(含未知地区)", "sales_share_of_month_pct", "pct"),
    _c("占已知地区销量(不含未知)", "sales_share_of_known_region_pct", "pct"),
    _c("地区内占比", "share_within_region_pct", "pct"),
    _c("占当月ASIN(含未知地区)", "asin_share_of_month_pct", "pct"),
]
_SEASONALITY = [
    _c("年份", "year", "int"),
    _c("已观测月数", "months_observed", "int"),
    _Col("12月齐全", lambda y: y.is_complete, "bool"),
    _Col("缺失月份", lambda y: "、".join(y.missing_months)),
    _c("高月", "high_month"),
    _c("高月销量", "high_sales", "int"),
    _c("低月", "low_month"),
    _c("低月销量", "low_sales", "int"),
    _c("覆盖提示", "coverage_warning"),
    _c("待验证假设", "cause_hypothesis"),
]
_SEGMENT = [
    _c("月份", "month"),
    _c("分类", "segment"),
    _c("销量", "sales", "int"),
    _c("ASIN数", "asin_count", "int"),
    _c("销量份额", "sales_share_pct", "pct"),
    _c("ASIN份额", "asin_share_pct", "pct"),
    _c("环比", "mom_pct", "pct"),
    _c("同比", "yoy_pct", "pct"),
]
_DECOMPOSITION = [
    _c("分类", "segment"),
    _c("基准月", "base_month"),
    _c("当前月", "current_month"),
    _c("基准市场销量", "base_market_sales", "int"),
    _c("当前市场销量", "current_market_sales", "int"),
    _c("基准分类销量", "base_segment_sales", "int"),
    _c("当前分类销量", "current_segment_sales", "int"),
    _c("市场变化", "market_change_pct", "pct"),
    _c("分类变化", "segment_change_pct", "pct"),
    _c("份额变化(pp)", "share_change_pp", "num"),
    _c("总变化销量", "total_change_sales", "num"),
    _c("市场贡献销量", "market_contribution_sales", "num"),
    _c("份额贡献销量", "share_contribution_sales", "num"),
]
_PERIOD = [
    _c("范围", "scope"),
    _c("对象", "segment"),
    _c("基准期", "base_period"),
    _c("当前期", "current_period"),
    _c("基准销量", "base_sales", "int"),
    _c("当前销量", "current_sales", "int"),
    _c("变化", "change_pct", "pct"),
    _c("基准份额", "base_share_pct", "pct"),
    _c("当前份额", "current_share_pct", "pct"),
    _c("份额变化(pp)", "share_change_pp", "num"),
]
_CONCENTRATION = [
    _c("月份", "month"),
    _c("总销量", "total_sales", "int"),
    _c("ASIN数", "asin_count", "int"),
    _c("Top3 ASIN占比", "top3_asin_share_pct", "pct"),
    _c("Top10 ASIN占比", "top10_asin_share_pct", "pct"),
    _c("Top3 品牌占比", "top3_brand_share_pct", "pct"),
    _c("Top10 品牌占比", "top10_brand_share_pct", "pct"),
]
_SELLER_CONCENTRATION = [
    _c("月份", "month"),
    _c("总销量", "total_sales", "int"),
    _c("ASIN数", "asin_count", "int"),
    _c("Top3 ASIN占比[口径=ASIN]", "top3_asin_share_pct", "pct"),
    _c("Top10 ASIN占比[口径=ASIN]", "top10_asin_share_pct", "pct"),
    _c("Top3 品牌占比[口径=品牌]", "top3_brand_share_pct", "pct"),
    _c("Top10 品牌占比[口径=品牌]", "top10_brand_share_pct", "pct"),
    _c("Top3 卖家占比[口径=seller_id]", "top3_seller_share_pct", "pct"),
    _c("Top10 卖家占比[口径=seller_id]", "top10_seller_share_pct", "pct"),
    _c("未知卖家销量(缺seller_id)", "unknown_seller_sales", "int"),
    _c("未知卖家ASIN数", "unknown_seller_asin_count", "int"),
    _c("未知卖家销量占比", "unknown_seller_share_pct", "pct"),
]
_POINT_CHANGE = [
    _c("名次", "rank", "int"),
    _c("基准月", "base_month"),
    _c("当前月", "current_month"),
    _c("基准销量", "base_point_sales", "int"),
    _c("当前销量", "current_point_sales", "int"),
    _c("绝对变化", "abs_change", "int"),
    _c("相对变化", "rel_change_pct", "pct"),
    _Col("方向", lambda r: sections._DIRECTION.get(r.direction, r.direction)),
    _c("口径(basis)", "basis"),
]
_ITERATION = [
    _c("顺序", "order", "int"),
    _c("分类", "segment"),
    _c("earliest_launch_date(上架日期)", "earliest_launch_date"),
    _c("最早上架ASIN", "earliest_launch_asin"),
    _c("earliest_observed_in_data(数据首次出现月,非上架日期)", "earliest_observed_in_data"),
    _c("ASIN数", "asin_count", "int"),
    _c("有效上架日期ASIN数", "valid_launch_date_asin_count", "int"),
    _c("排除:缺失", "excluded_missing", "int"),
    _c("排除:无法解析", "excluded_unparseable", "int"),
    _c("排除:未来日期", "excluded_future", "int"),
    _c("排除:过早日期", "excluded_too_early", "int"),
]
_RANK_LAYER = [
    _c("月份", "month"),
    _c("名次截止", "cutoff", "int"),
    _c("销量", "sales", "int"),
    _c("ASIN数", "asin_count", "int"),
    _c("销量份额", "sales_share_pct", "pct"),
    _c("同比", "yoy_pct", "pct"),
    _c("口径(basis)", "basis"),
]
_RANK_BAND = [
    _c("月份", "month"),
    _c("分层", "band"),
    _c("起始名次", "start_rank", "int"),
    _c("结束名次", "end_rank", "int"),
    _c("销量", "sales", "int"),
    _c("ASIN数", "asin_count", "int"),
    _c("销量份额", "sales_share_pct", "pct"),
    _c("同比", "yoy_pct", "pct"),
    _c("口径(basis)", "basis"),
]
_HEAD_SUCTION = [
    _c("月份", "month"),
    _c("基准月", "base_month"),
    _c("头部截止", "head_cutoff", "int"),
    _c("宽口径截止", "broad_cutoff", "int"),
    _c("头部同比", "head_yoy_pct", "pct"),
    _c("宽口径同比", "broad_yoy_pct", "pct"),
    _c("头部减宽口径(pp)", "head_minus_broad_pp", "num"),
    _c("模式", "pattern"),
    _c("口径(basis)", "basis"),
]
_NEW_PRODUCT = [
    _c("ASIN", "asin"),
    _c("锚定月", "anchor_month"),
    _c("首次观察月", "first_observed_month"),
    _c("首次达标月", "first_threshold_month"),
    _c("首次达标月序", "first_threshold_month_index", "int"),
    _c("是否成功", "launch_success", "bool"),
    _c("成功月", "success_month"),
    _c("当前月", "current_month"),
    _c("当前销量", "current_sales", "int"),
    _c("当前是否观察到", "current_observed", "bool"),
    _c("当前是否存活", "current_survival", "bool"),
    _c("卖家所属地", "seller_region"),
    _c("账号类型", "account_type"),
    _c("分类", "segment"),
]


def write_research_workbook(
    input: ResearchDossierInput,
    output_path: Path,
    *,
    generated_at: str | None = None,
) -> Path:
    """把 dossier 内容写成可离线打开、可编辑的 .xlsx。

    同输入同输出:工作簿内不写当前时间,`generated_at` 由调用方注入;
    zip 条目时间戳与文档属性固定,字节级可复现。
    """
    openpyxl = _import_openpyxl()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "research_workflow"
    wb.properties.created = _FIXED_DATETIME
    wb.properties.modified = _FIXED_DATETIME

    _readme_sheet(wb, input, generated_at)
    _records_sheet(wb, input)
    _classification_sheet(wb, input)
    metrics = input.market_metrics
    _table_sheet(wb, "市场_月度", _MARKET, metrics.market)
    _table_sheet(wb, "市场_地区", _REGION, metrics.seller_regions)
    _table_sheet(wb, "市场_分类", _SEGMENT, metrics.segments)
    _table_sheet(wb, "市场_分解", _DECOMPOSITION, metrics.decompositions)
    _table_sheet(wb, "市场_跨期", _PERIOD, metrics.period_comparisons or [])
    _table_sheet(wb, "市场_国别分类透视", _REGION_SEGMENT, metrics.region_segments or None)
    _seasonality_sheet(wb, input)
    comp = input.competition_metrics
    _table_sheet(wb, "竞争_集中度", _CONCENTRATION, comp.concentration if comp else None)
    _table_sheet(wb, "竞争_卖家集中度", _SELLER_CONCENTRATION, comp.concentration if comp else None)
    _table_sheet(wb, "竞争_BSR层级", _RANK_LAYER, comp.rank_layers if comp else None)
    _table_sheet(wb, "竞争_头腰尾", _RANK_BAND, comp.rank_bands if comp else None)
    _table_sheet(wb, "竞争_虹吸", _HEAD_SUCTION, comp.head_suction if comp else None)
    _point_rank_matrix_sheet(wb, input)
    _table_sheet(
        wb, "竞争_单点名次变化", _POINT_CHANGE,
        (comp.point_rank_changes or None) if comp else None,
    )
    new = input.new_product_analysis
    _table_sheet(wb, "新品", _NEW_PRODUCT, new.outcomes if new else None)
    _retention_sheet(wb, input)
    _table_sheet(wb, "产品迭代", _ITERATION, input.iteration.segments if input.iteration else None)
    _iteration_notes(wb, input)
    _quality_sheet(wb, input)
    _missing_sheet(wb, input)
    _sources_sheet(wb, input)

    buffer = BytesIO()
    wb.save(buffer)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(_normalize_zip(buffer.getvalue()))
    return output


# ---------------------------------------------------------------- sheets


def _readme_sheet(wb: Any, input: ResearchDossierInput, generated_at: str | None) -> None:
    ws = wb.create_sheet("说明")
    classification = input.classification
    approvals = input.approval_history
    rows: list[tuple[str, object]] = [
        ("项目", input.project_id),
        ("产品", input.product),
        ("站点", input.marketplace),
        ("数据版本", input.data_version or NOT_PROVIDED),
        ("数据哈希", input.data_hash or NOT_PROVIDED),
        ("分类版本", classification.version if classification else NOT_PROVIDED),
        ("生成时间", generated_at or NOT_PROVIDED),
        ("审批历史", _approval_summary(approvals) if approvals else NOT_PROVIDED),
    ]
    for r, (label, value) in enumerate(rows, start=1):
        _put(ws, r, 1, label, bold=True)
        _put(ws, r, 2, value)
    start = len(rows) + 2
    _put(ws, start, 1, "人工决策栏", bold=True)
    _put(ws, start, 2, DECISION_NOTE)
    for offset, label in enumerate(("最终结论", "决策人", "决策日期", "决策原因"), start=1):
        _put(ws, start + offset, 1, label, bold=True)  # B 列刻意留空,由人工填写
    scope_start = start + 6
    _put(ws, scope_start, 1, sections.SCOPE_TITLE, bold=True)
    for offset, note in enumerate(sections.SCOPE_NOTES, start=1):
        _put(ws, scope_start + offset, 1, f"口径{offset}", bold=True)
        _put(ws, scope_start + offset, 2, note)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 70


def _seasonality_sheet(wb: Any, input: ResearchDossierInput) -> None:
    result = input.market_metrics.seasonality
    ws = _table_sheet(wb, "市场_季节性", _SEASONALITY, result.years or None)
    if not result.years:
        return
    footer = [
        ("完整年份", "、".join(str(y) for y in result.complete_years) or "-"),
        ("跨完整年份重复出现的高月(月份序号)", result.recurring_high_month or "-"),
        ("跨完整年份重复出现的低月(月份序号)", result.recurring_low_month or "-"),
        ("整体覆盖提示", result.coverage_warning or "-"),
        ("口径", "仅 12 个月齐全的年份判断高低月;缺月不当 0;“待验证假设”不是已验证原因"),
    ]
    start = len(result.years) + 3
    for offset, (label, value) in enumerate(footer):
        _put(ws, start + offset, 1, label, bold=True)
        _put(ws, start + offset, 2, value)


def _point_rank_matrix_sheet(wb: Any, input: ResearchDossierInput) -> None:
    comp = input.competition_metrics
    matrix = comp.point_rank_matrix if comp else None
    if matrix is None or not matrix.rows:
        _table_sheet(wb, "竞争_单点名次矩阵", [], None)
        return
    columns = [
        _c("月份", "month"),
        _c("排序口径", "ordering_basis"),
        _c("口径(basis)", "basis"),
    ] + [
        _Col(f"{rank}名", lambda row, key=str(rank): row.point_sales.get(key), "int")
        for rank in matrix.ranks
    ]
    ws = _table_sheet(wb, "竞争_单点名次矩阵", columns, sorted(matrix.rows, key=lambda r: r.month))
    note_row = len(matrix.rows) + 3
    _put(ws, note_row, 1, "说明", bold=True)
    _put(ws, note_row, 2, "每格为该名次单个 ASIN 的月销量(basis=point,非累计);ASIN 数不足该名次时留空,不补 0")


def _retention_sheet(wb: Any, input: ResearchDossierInput) -> None:
    rows = sections.retention_rows(input)
    if not rows:
        _table_sheet(wb, "新品_留存", [], None)
        return
    ws = wb.create_sheet("新品_留存")
    for col, header in enumerate(("指标", "值", "分母与口径"), start=1):
        _put(ws, 1, col, header, bold=True)
    for r, (label, value, kind, note) in enumerate(rows, start=2):
        _put(ws, r, 1, label)
        _put(ws, r, 2, value, kind="text" if kind == "text" else kind)
        _put(ws, r, 3, note)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 70
    ws.freeze_panes = "A2"


def _iteration_notes(wb: Any, input: ResearchDossierInput) -> None:
    if input.iteration is None:
        return
    ws = wb["产品迭代"]
    start = len(input.iteration.segments) + 3
    _put(ws, start, 1, "说明", bold=True)
    for offset, note in enumerate(input.iteration.notes, start=1):
        _put(ws, start + offset, 1, note)


def _approval_summary(history: Sequence[dict[str, object]]) -> str:
    lines = []
    for item in history:
        lines.append("; ".join(f"{k}={item[k]}" for k in sorted(item)))
    return "\n".join(lines)


def _records_sheet(wb: Any, input: ResearchDossierInput) -> None:
    records = input.records
    ranks = _cleaned_ranks(records or [])
    columns = [
        _Col("ASIN", lambda r: r.asin),
        _Col("品牌", lambda r: r.brand),
        _Col("标题", lambda r: r.title),
        _Col("月销量", lambda r: r.monthly_sales, "int"),
        _Col("月份", lambda r: r.month),
        _Col("节点排名", lambda r: r.node_rank, "int"),
        _Col("清洗后名次", lambda r: r.sales_rank if r.sales_rank is not None else ranks[id(r)], "int"),
        _Col("卖家所属地", lambda r: r.seller_region),
        _Col("上架日期", lambda r: r.launch_date),
        _Col("父ASIN", lambda r: r.parent_asin),
        _Col("图片URL", lambda r: r.image_url),
    ]
    _table_sheet(wb, "表1_清洗月度数据", columns, records)


def _cleaned_ranks(records: Sequence[Any]) -> dict[int, int]:
    """月内按销量降序的名次;销量相同依次比节点排名、ASIN,保证确定性。"""
    by_month: dict[str, list[Any]] = {}
    for record in records:
        by_month.setdefault(record.month, []).append(record)
    ranks: dict[int, int] = {}
    for items in by_month.values():
        items.sort(
            key=lambda r: (
                -r.monthly_sales,
                r.node_rank if r.node_rank is not None else float("inf"),
                r.asin,
            )
        )
        for position, record in enumerate(items, start=1):
            ranks[id(record)] = position
    return ranks


def _classification_sheet(wb: Any, input: ResearchDossierInput) -> None:
    batch = input.classification
    if batch is None:
        _table_sheet(wb, "表3_分类结果", [], None)
        return
    edited = set(batch.human_edits)
    columns = [
        _Col("ASIN", lambda r: r.asin),
        _Col("标签", lambda r: r.label),
        _Col("依据", lambda r: r.rationale),
        _Col("是否人工修改", lambda r: "是" if r.asin in edited or r.source == "human" else "否"),
        _Col("置信度", lambda r: r.confidence),
        _Col("来源", lambda r: r.source),
    ]
    _table_sheet(wb, "表3_分类结果", columns, [batch.rows[k] for k in sorted(batch.rows)])


def _quality_sheet(wb: Any, input: ResearchDossierInput) -> None:
    rows = list(input.data_quality)
    columns = [
        _Col("指标", lambda r: r.get("label", "")),
        _Col("值", lambda r: r.get("value")),
    ]
    _table_sheet(wb, "数据质量", columns, rows or None)


def _missing_sheet(wb: Any, input: ResearchDossierInput) -> None:
    columns = [_Col("缺失证据", lambda r: r)]
    ws = _table_sheet(wb, "缺失证据", columns, list(input.missing_evidence))
    if not input.missing_evidence:
        ws["A1"].value = "无已登记缺口"


def _sources_sheet(wb: Any, input: ResearchDossierInput) -> None:
    columns = [
        _Col("ID", lambda s: s.artifact_id),
        _Col("名称", lambda s: s.name),
        _Col("相对路径", lambda s: s.relative_path or s.name),
        _Col("SHA256", lambda s: s.sha256),
        _Col("类型", lambda s: s.kind),
        _Col("来源", lambda s: s.source),
        _Col("说明", lambda s: s.note),
    ]
    _table_sheet(wb, "来源与哈希", columns, list(input.source_artifacts) or None)


# ---------------------------------------------------------------- helpers


def _table_sheet(
    wb: Any,
    title: str,
    columns: Sequence[_Col],
    rows: Iterable[Any] | None,
) -> Any:
    """rows 为 None 表示本次运行未提供该数据,整张表只写一行说明。"""
    ws = wb.create_sheet(title)
    values = None if rows is None else list(rows)
    if values is None:
        _put(ws, 1, 1, NOT_PROVIDED)
        return ws
    for col_idx, col in enumerate(columns, start=1):
        _put(ws, 1, col_idx, col.header, bold=True)
        ws.column_dimensions[_letter(col_idx)].width = max(12, len(col.header) * 2 + 2)
    for row_idx, row in enumerate(values, start=2):
        for col_idx, col in enumerate(columns, start=1):
            _put(ws, row_idx, col_idx, col.get(row), kind=col.kind)
    if columns:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{_letter(len(columns))}{max(1, len(values) + 1)}"
    return ws


def _letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)


def _put(ws: Any, row: int, col: int, value: object, *, kind: str = "text", bold: bool = False) -> None:
    cell = ws.cell(row=row, column=col)
    if value is None:
        pass
    elif kind == "pct":
        cell.value = float(value) / 100.0  # type: ignore[arg-type]
        cell.number_format = "0.00%"
    elif kind == "int":
        cell.value = int(value)  # type: ignore[call-overload]
    elif kind == "num":
        cell.value = float(value)  # type: ignore[arg-type]
        cell.number_format = "#,##0.00"
    elif kind == "bool":
        cell.value = bool(value)
    else:
        text = value if isinstance(value, (int, float, bool)) else str(value)
        cell.value = text
        if isinstance(text, str):
            cell.data_type = "s"  # 防止以 = 开头的标题被当作公式
    if bold:
        from openpyxl.styles import Font

        cell.font = Font(bold=True)


def _normalize_zip(raw: bytes) -> bytes:
    """重写 zip 条目时间戳,使同输入得到字节相同的 xlsx。"""
    out = BytesIO()
    with ZipFile(BytesIO(raw)) as src, ZipFile(out, "w", ZIP_DEFLATED) as dst:
        for name in src.namelist():
            info = ZipInfo(name, date_time=_FIXED_ZIP_TIME)
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dst.writestr(info, src.read(name))
    return out.getvalue()
