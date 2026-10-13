"""三张过程表导出/上传、差异留痕，以及综合 Excel 成品。"""
from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from scripts.run_research import main


@pytest.fixture()
def cli(tmp_path, capsys):
    def run(*argv):
        capsys.readouterr()
        main(["--db", str(tmp_path / "db.sqlite"), "--root", str(tmp_path / "projects"), *argv])
        return json.loads(capsys.readouterr().out)

    def fail(*argv):
        capsys.readouterr()
        with pytest.raises(SystemExit):
            main(["--db", str(tmp_path / "db.sqlite"), "--root", str(tmp_path / "projects"), *argv])
        return json.loads(capsys.readouterr().out)["error"]

    run.fail, run.tmp = fail, tmp_path
    return run


def _w(tmp_path, name, text):
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return str(p)


CSV = "ASIN,月销量,品牌,商品标题,小类BSR,卖家所属地,评分,评分数\nN1,100,NB,Nugget Maker,1,CN,4.5,100\nB1,50,BB,Bullet Maker,2,US,4.1,40\n"


def _start(cli):
    pid = cli("init", "--product", "Ice Maker", "--marketplace", "US")["project_id"]
    out = cli("ingest", "--project", pid, "--file", _w(cli.tmp, "2026-01.csv", CSV))
    return pid, out["checkpoint_id"]


def _edit_b1_sales(xlsx, value):
    wb = load_workbook(xlsx)
    ws = wb["表1"]
    headers = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[headers.index("ASIN")].value == "B1":
            row[headers.index("月销量")].value = value
    wb.save(xlsx)


def test_table1_export_edit_import_creates_audit_and_new_checkpoint(cli):
    pid, cp = _start(cli)
    xlsx = cli.tmp / "t1.xlsx"
    assert cli("table-export", "--project", pid, "--table", "1", "--output", str(xlsx))["rows"] == 2

    # 未改动 => 拒绝，不制造空版本
    assert "完全一致" in cli.fail("table-import", "--project", pid, "--table", "1", "--file", str(xlsx))

    _edit_b1_sales(xlsx, 70)
    out = cli("table-import", "--project", pid, "--table", "1", "--file", str(xlsx))
    assert out["checkpoint_step"] == "data_review" and out["checkpoint_id"] != cp
    pdir = cli.tmp / "projects" / pid
    recs = json.loads((pdir / "records.json").read_text(encoding="utf-8"))["2026-01"]
    assert {r["asin"]: r["monthly_sales"] for r in recs} == {"N1": 100, "B1": 70}
    log = json.loads((pdir / "table_edits.json").read_text(encoding="utf-8"))
    assert log[0]["changes"] == [{"month": "2026-01", "asin": "B1", "field": "monthly_sales", "from": 50, "to": 70}]
    sources = json.loads((pdir / "sources.json").read_text(encoding="utf-8"))
    assert any(v["kind"] == "human_table" for v in sources.values())
    # 旧检查点已失效，只能裁决新的
    assert "过期" in cli.fail("resolve", "--project", pid, "--checkpoint", cp, "--approve")
    cli("resolve", "--project", pid, "--checkpoint", out["checkpoint_id"], "--approve")


def test_table1_import_rejects_missing_sales_and_month_changes(cli):
    pid, _ = _start(cli)
    xlsx = cli.tmp / "t1.xlsx"
    cli("table-export", "--project", pid, "--table", "1", "--output", str(xlsx))
    _edit_b1_sales(xlsx, None)
    assert "不能视为 0" in cli.fail("table-import", "--project", pid, "--table", "1", "--file", str(xlsx))
    csv = cli.tmp / "t1.csv"
    cli("table-export", "--project", pid, "--table", "1", "--output", str(csv))
    text = csv.read_text(encoding="utf-8-sig").replace("2026-01", "2026-02")
    csv.write_text(text, encoding="utf-8-sig")
    assert "月份集合" in cli.fail("table-import", "--project", pid, "--table", "1", "--file", str(csv))


def test_table3_edits_file_only_applies_changed_rows(cli):
    pid, cp = _start(cli)
    cli("resolve", "--project", pid, "--checkpoint", cp, "--approve")
    out = cli("classify", "--project", pid, "--label", "nugget=咀嚼冰", "--label", "bullet=子弹冰",
              "--rule", "nugget=nugget")
    t3 = cli.tmp / "t3.csv"
    cli("table-export", "--project", pid, "--table", "3", "--output", str(t3))
    text = t3.read_text(encoding="utf-8-sig")
    assert "B1" in text and "unknown" in text
    lines = text.splitlines()
    header = lines[0].split(",")
    i_asin, i_cls = header.index("ASIN"), header.index("分类")
    new = [lines[0]]
    for line in lines[1:]:
        cells = line.split(",")
        if cells[i_asin] == "B1":
            cells[i_cls] = "bullet"
        new.append(",".join(cells))
    t3.write_text("\n".join(new) + "\n", encoding="utf-8-sig")
    out = cli("resolve", "--project", pid, "--checkpoint", out["checkpoint_id"], "--approve", "--edits-file", str(t3))
    assert out["edits"] == {"B1": "bullet"}


def test_dossier_includes_editable_workbook_with_blank_human_decision(cli):
    pid, cp = _start(cli)
    cli("resolve", "--project", pid, "--checkpoint", cp, "--approve")
    out = cli("classify", "--project", pid, "--label", "nugget=咀嚼冰", "--label", "bullet=子弹冰", "--rule", "nugget=nugget")
    cli("resolve", "--project", pid, "--checkpoint", out["checkpoint_id"], "--approve", "--edit", "B1=bullet")
    out = cli("analyze", "--project", pid)
    cli("resolve", "--project", pid, "--checkpoint", out["checkpoint_id"], "--approve", "--select", "")
    dossier = cli.tmp / "out"
    result = cli("dossier", "--project", pid, "--output", str(dossier))
    assert "workbook.xlsx" in result["files"]
    wb = load_workbook(dossier / "workbook.xlsx")
    assert {"说明", "原始数据", "表1_清洗Top", "表3_分类", "市场月度", "集中度", "逐名次销量"} <= set(wb.sheetnames)
    notes = {r[0].value: r[1].value for r in wb["说明"].iter_rows(min_row=2)}
    assert notes["人工决策（Yes/No）"] in (None, "")
    manifest = json.loads((dossier / "manifest.json").read_text(encoding="utf-8"))
    assert "workbook.xlsx" in {f["path"] for f in manifest["files"]}
    raw_sheet = wb["原始数据"]
    assert raw_sheet.max_row == 3  # 表头 + 两条原始候选行
